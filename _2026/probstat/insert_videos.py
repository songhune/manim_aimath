# -*- coding: utf-8 -*-
"""강의 슬라이드에 manim 영상을 넣은 사본을 만든다. 원본 pptx 는 건드리지 않는다.

씬을 다시 렌더한 뒤 이 스크립트를 그대로 돌리면 슬라이드 구성은 유지한 채
영상만 새것으로 바뀐다. 손으로 한 장씩 넣지 않는다.

    python _2026/probstat/insert_videos.py                       # 전부
    python _2026/probstat/insert_videos.py PS1_02_restyled.pptx   # 이 원본의 job 만

만드는 것
    [0901]오리엔테이션.pptx      → [0901]오리엔테이션_영상.pptx
    PS1_01_restyled.pptx         → PS1_01_restyled_영상.pptx
    PS1_01_restyled_한글.pptx    → PS1_01_restyled_한글_영상.pptx
    PS1_02_restyled.pptx         → PS1_02_restyled_영상.pptx
    PS1_02_restyled_한글.pptx    → PS1_02_restyled_한글_영상.pptx

영상은 16:9 전체 화면 슬라이드로 넣는다. 본문 슬라이드 한구석에 작게 넣으면
1080p 영상의 축 이름과 자막이 강의실 스크린에서 읽히지 않는다.
재생은 PowerPoint 기본값인 '클릭할 때'이며, 교안의 진행 순서(말 → 판서 → 영상)와 맞는다.
"""
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "_harness"))
from fix_deck_fonts import fix as fix_fonts   # 하네스 3.5. 덱을 쓰기 전에 옛 글꼴 이름을 고친다

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
BLANK_LAYOUT = 6                       # 슬라이드 마스터의 '빈 화면'

# 이 파일: 2026-2/manim_songhune/_2026/probstat/insert_videos.py
SEMESTER = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
DECKS = os.path.join(SEMESTER, "확률과통계", "확통 자료", "수업안")
VIDEOS = os.path.join(DECKS, "영상")

# 이미 영상 슬라이드인 것은 파일만 갈아 끼우고, 나머지는 새 슬라이드로 끼워 넣는다.
#
# insert=(N, [Scene ...]) 은 "원본 N번 슬라이드 뒤" 라는 뜻이다.
# anchors 는 그 영상이 다루는 개념 슬라이드 번호다.
#
# place 가 배치 규칙을 정한다 (하네스 3.6).
#   "before" — 영상이 개념 슬라이드 바로 앞에 온다. 지금의 규칙이다.
#              예제 풀이 영상은 문제 슬라이드 다음, 풀이 슬라이드 앞에 놓인다.
#   "after"  — 영상이 개념 슬라이드 바로 뒤에 온다. 2026-09-04 이전 규칙이다.
#              1차·2차에 이미 나간 오리엔테이션·PS1_01 덱은 교안의 슬라이드 표와 판서 덱이
#              그 순서에 맞춰져 있으므로 그대로 둔다.
# verify() 가 만들어진 파일에서 실제 순서를 다시 센다.
JOBS = [
    dict(
        src="[0901]오리엔테이션.pptx",
        dst="[0901]오리엔테이션_영상.pptx",
        place="after",
        replace={5: "MontyHall", 13: "PopulationAndSample"},
        # 4 = 강의개요(몬티 홀), 6 = 확률의 두 관점, 12 = Ch1 도입, 14 = 모집단과 표본
        insert=[(6, ["FrequencyView", "BeliefUpdate"]),
                (12, ["WhyVariability"]),
                (14, ["ProbabilityVsInference"])],
        move=[(13, 14)],               # 영상이 본문 슬라이드 뒤에 오게 한다
        anchors={"MontyHall": 4, "FrequencyView": 6, "BeliefUpdate": 6,
                 "WhyVariability": 12, "PopulationAndSample": 14,
                 "ProbabilityVsInference": 14},
    ),
    dict(
        src="PS1_01_restyled.pptx",
        dst="PS1_01_restyled_영상.pptx",
        place="after",
        insert=[(4, ["WhyVariability"]),          # 4 = Variability in Scientific Data
                (5, ["ProbabilityVsInference"]),  # 5 = The Role of Probability
                (6, ["BiasedSample"]),            # 6 = Collection of data
                (8, ["MeanVsMedian"]),            # 7 = Definition 1.1, 8 = 평균·중앙값 그림
                (9, ["VarianceFormula"]),         # 9 = Measures of Variability (정의 1.3)
                (10, ["VarianceAsSquares"]),      # 10 = 분산 예제
                (12, ["BarVsHistogram", "HistogramFromTable"]),   # 11 = 막대, 12 = 히스토그램
                (13, ["Skewness"])],              # 13 = 히스토그램 그림
        anchors={"WhyVariability": 4, "ProbabilityVsInference": 5, "BiasedSample": 6,
                 "MeanVsMedian": 8, "VarianceFormula": 9, "VarianceAsSquares": 10,
                 "BarVsHistogram": 12, "HistogramFromTable": 12, "Skewness": 13},
    ),
    dict(
        # 한글판은 영문판과 장 수·순서가 같으므로 삽입 위치를 그대로 쓴다.
        src="PS1_01_restyled_한글.pptx",
        dst="PS1_01_restyled_한글_영상.pptx",
        place="after",
        insert=[(4, ["WhyVariability"]), (5, ["ProbabilityVsInference"]),
                (6, ["BiasedSample"]), (8, ["MeanVsMedian"]),
                (9, ["VarianceFormula"]), (10, ["VarianceAsSquares"]),
                (12, ["BarVsHistogram", "HistogramFromTable"]),
                (13, ["Skewness"])],
        anchors={"WhyVariability": 4, "ProbabilityVsInference": 5, "BiasedSample": 6,
                 "MeanVsMedian": 8, "VarianceFormula": 9, "VarianceAsSquares": 10,
                 "BarVsHistogram": 12, "HistogramFromTable": 12, "Skewness": 13},
    ),
    dict(
        # PS1_02 덱은 세 벌이다: 영문 영상판(교수자용) · 영문 학생본(도구/학생배포본_생성.py) · 한글 최종본.
        # 개념 영상은 개념 슬라이드 바로 앞, 예제 영상은 문제 슬라이드 다음(풀이 슬라이드 앞)이다.
        # 3차(9/8) 는 2.1–2.3 = 원본 1–26, 4차(9/11) 는 2.4–2.5 = 원본 27–41.
        src="PS1_02_restyled.pptx",
        dst="PS1_02_restyled_영상.pptx",
        place="before",
        drop=[47],                                       # 원본 47 은 글이 없는 빈 슬라이드다. 뺀다 (2026-09-10)
        extra=[(45, "suneung_cond")],                    # 45 뒤에 수능형 연습 문제 슬라이드 (2026-09-10)
        insert=[(3, ["SampleSpace"]),                        # 4 = Definition 2.1 · Example 2.1
                (7, ["EventsAndSetOps"]),                    # 8 = Definition 2.2–2.6
                (11, ["MultiplicationRule"]),                # 12 = Rule 2.1 · 2.2
                (13, ["Example213Dice"]),                    # 13 = Ex 2.13 문제, 14 = 풀이
                (15, ["Example215Club", "Example217EvenNumbers"]),   # 15 = Ex 2.15·2.16·2.17 문제, 16 = 풀이
                (16, ["Permutations"]),                      # 17 = Definition 2.7 · Theorem 2.1 · 2.2
                (17, ["Combinations"]),                      # 18 = Theorem 2.3–2.6
                (19, ["Example218Awards"]),                  # 19 = Ex 2.18 문제, 20 = 풀이
                (25, ["Example222Cartridges"]),              # 25 = Ex 2.22·2.23 문제, 26 = 풀이
                (26, ["ProbabilityOfEvent"]),                # 27 = Definition 2.9 · Rule 2.3
                (28, ["Example225LoadedDie"]),               # 28 = Ex 2.24·2.25 문제, 29 = 풀이
                (32, ["Example228Poker"]),                   # 32 = Ex 2.28 문제, 33 = 풀이
                (34, ["AdditionRule"]),                      # 35 = Theorem 2.7 · 2.9
                (36, ["Example229Jobs"]),                    # 36 = Ex 2.29·2.30 문제, 37 = 풀이
                (40, ["Example233Cable"]),                   # 40 = Ex 2.33 문제, 41 = 풀이
                # 2.6–2.8 (week03.py). 2026-09-08 추가.
                (41, ["ConditionalProbability"]),            # 42 = Definition 2.10
                (44, ["Example234Flights"]),                 # 44 = Ex 2.34·2.35 문제, 45 = 풀이
                (45, ["SuneungConditional", "Independence"]),   # 45 뒤: 수능형 문제 슬라이드(extra) → 그 풀이 영상 → 46 = Definition 2.11 앞 개념 영상
                (47, ["ProductRule"]),                       # 48 = Theorem 2.10–2.12
                (49, ["Example236Fuses"]),                   # 49 = Ex 2.36·2.37 문제, 50 = 풀이
                (51, ["Example238Emergency"]),               # 51 = Ex 2.38·2.40 문제, 52 = 풀이
                (52, ["TotalProbability"]),                  # 53 = Theorem 2.13
                (54, ["Example241Machines"]),                # 54 = Ex 2.41 문제, 55 = 풀이
                (55, ["BayesRule"]),                         # 56 = Theorem 2.14
                (57, ["Example242Bayes"])],                  # 57 = Ex 2.42·2.43 문제, 58 = 풀이
        anchors={"SampleSpace": 4, "EventsAndSetOps": 8, "MultiplicationRule": 12,
                 "Example213Dice": 14, "Example215Club": 16, "Example217EvenNumbers": 16,
                 "Permutations": 17, "Combinations": 18, "Example218Awards": 20,
                 "Example222Cartridges": 26, "ProbabilityOfEvent": 27, "Example225LoadedDie": 29,
                 "Example228Poker": 33, "AdditionRule": 35, "Example229Jobs": 37,
                 "Example233Cable": 41,
                 "ConditionalProbability": 42, "Example234Flights": 45, "SuneungConditional": 46, "Independence": 46,
                 "ProductRule": 48, "Example236Fuses": 50, "Example238Emergency": 52,
                 "TotalProbability": 53, "Example241Machines": 55, "BayesRule": 56,
                 "Example242Bayes": 58},
    ),
    dict(
        # 한글판은 영문판과 장 수·순서가 같으므로 삽입 위치를 그대로 쓴다.
        src="PS1_02_restyled_한글.pptx",
        dst="PS1_02_restyled_한글_영상.pptx",
        place="before",
        drop=[47],
        extra=[(45, "suneung_cond")],
        insert=[(3, ["SampleSpace"]),                        # 4 = Definition 2.1 · Example 2.1
                (7, ["EventsAndSetOps"]),                    # 8 = Definition 2.2–2.6
                (11, ["MultiplicationRule"]),                # 12 = Rule 2.1 · 2.2
                (13, ["Example213Dice"]),                    # 13 = Ex 2.13 문제, 14 = 풀이
                (15, ["Example215Club", "Example217EvenNumbers"]),   # 15 = Ex 2.15·2.16·2.17 문제, 16 = 풀이
                (16, ["Permutations"]),                      # 17 = Definition 2.7 · Theorem 2.1 · 2.2
                (17, ["Combinations"]),                      # 18 = Theorem 2.3–2.6
                (19, ["Example218Awards"]),                  # 19 = Ex 2.18 문제, 20 = 풀이
                (25, ["Example222Cartridges"]),              # 25 = Ex 2.22·2.23 문제, 26 = 풀이
                (26, ["ProbabilityOfEvent"]),                # 27 = Definition 2.9 · Rule 2.3
                (28, ["Example225LoadedDie"]),               # 28 = Ex 2.24·2.25 문제, 29 = 풀이
                (32, ["Example228Poker"]),                   # 32 = Ex 2.28 문제, 33 = 풀이
                (34, ["AdditionRule"]),                      # 35 = Theorem 2.7 · 2.9
                (36, ["Example229Jobs"]),                    # 36 = Ex 2.29·2.30 문제, 37 = 풀이
                (40, ["Example233Cable"]),                   # 40 = Ex 2.33 문제, 41 = 풀이
                # 2.6–2.8 (week03.py). 2026-09-08 추가.
                (41, ["ConditionalProbability"]),            # 42 = Definition 2.10
                (44, ["Example234Flights"]),                 # 44 = Ex 2.34·2.35 문제, 45 = 풀이
                (45, ["SuneungConditional", "Independence"]),   # 45 뒤: 수능형 문제 슬라이드(extra) → 그 풀이 영상 → 46 = Definition 2.11 앞 개념 영상
                (47, ["ProductRule"]),                       # 48 = Theorem 2.10–2.12
                (49, ["Example236Fuses"]),                   # 49 = Ex 2.36·2.37 문제, 50 = 풀이
                (51, ["Example238Emergency"]),               # 51 = Ex 2.38·2.40 문제, 52 = 풀이
                (52, ["TotalProbability"]),                  # 53 = Theorem 2.13
                (54, ["Example241Machines"]),                # 54 = Ex 2.41 문제, 55 = 풀이
                (55, ["BayesRule"]),                         # 56 = Theorem 2.14
                (57, ["Example242Bayes"])],                  # 57 = Ex 2.42·2.43 문제, 58 = 풀이
        anchors={"SampleSpace": 4, "EventsAndSetOps": 8, "MultiplicationRule": 12,
                 "Example213Dice": 14, "Example215Club": 16, "Example217EvenNumbers": 16,
                 "Permutations": 17, "Combinations": 18, "Example218Awards": 20,
                 "Example222Cartridges": 26, "ProbabilityOfEvent": 27, "Example225LoadedDie": 29,
                 "Example228Poker": 33, "AdditionRule": 35, "Example229Jobs": 37,
                 "Example233Cable": 41,
                 "ConditionalProbability": 42, "Example234Flights": 45, "SuneungConditional": 46, "Independence": 46,
                 "ProductRule": 48, "Example236Fuses": 50, "Example238Emergency": 52,
                 "TotalProbability": 53, "Example241Machines": 55, "BayesRule": 56,
                 "Example242Bayes": 58},
    ),
]


def duration(path):
    out = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "csv=p=0", path], capture_output=True, text=True, check=True)
    return float(out.stdout.strip())


def poster(name, out_dir):
    """미리보기 이미지. 마지막 장면을 쓴다 — 모든 씬이 결론 식을 강조한 화면으로 끝나므로 그 장면이
    영상의 요지다(2026-09-10. 그전에는 중반 55% 지점이었다). 학생 배포본의 링크 슬라이드도 이 그림을 쓴다."""
    mp4 = os.path.join(VIDEOS, name + ".mp4")
    png = os.path.join(out_dir, name + ".png")
    if not os.path.exists(png):
        subprocess.run(
            ["ffmpeg", "-v", "error", "-ss", f"{max(duration(mp4) - 0.4, 0):.2f}",
             "-i", mp4, "-frames:v", "1", "-y", png], check=True)
    return mp4, png


# 덱에 끼우는 문제 슬라이드. 원본에 없는 문제(수능 기출)를 제목·본문(·표)으로 만든다. 2026-09-10.
# 문제 문장은 원문 그대로 옮기고 출처를 적는다. 2026학년도 수능 확률과 통계 28번(조건부확률 + 독립시행).
EXTRA_SLIDES = {
    "suneung_cond": dict(
        title="2.6 Conditional Probability · 2026학년도 수능 확률과 통계 28번",
        body=[
            (0, "16개의 공과 1부터 6까지의 자연수가 하나씩 적혀 있는 여섯 개의 빈 상자가 있다. 한 개의 주사위를 사용하여 다음 시행을 한다."),
            (1, "주사위를 한 번 던져 나온 눈의 수가 k일 때, k가 홀수이면 1, 3, 5가 적힌 상자에 공을 각각 1개씩 넣고, "
                "k가 짝수이면 k의 약수가 적힌 상자에 공을 각각 1개씩 넣는다."),
            (0, "이 시행을 4번 반복한 후 여섯 개의 상자에 들어 있는 모든 공의 개수의 합이 홀수일 때, "
                "3이 적힌 상자에 들어 있는 공의 개수가 2가 적힌 상자에 들어 있는 공의 개수보다 1개 더 많을 확률은?"),
            (0, "①  1/8        ②  3/16        ③  1/4        ④  5/16        ⑤  3/8"),
            (1, "출처: 2026학년도 대학수학능력시험 수학 영역 확률과 통계 28번 (한국교육과정평가원)"),
        ],
    ),
}


def extra_slide(prs, key):
    """제목 + 본문 + 표 슬라이드 하나를 덱 끝에 만든다(자리는 build 가 옮긴다)."""
    spec = EXTRA_SLIDES[key]
    slide = prs.slides.add_slide(prs.slide_layouts[1])          # 제목 및 내용
    slide.shapes.title.text = spec["title"]
    body = slide.placeholders[1]
    body.width, body.height = int(prs.slide_width * 0.86), int(prs.slide_height * 0.42)
    tf = body.text_frame
    tf.text = spec["body"][0][1]
    for level, text in spec["body"][1:]:
        para = tf.add_paragraph()
        para.text = text
        para.level = level
    if not spec.get("table"):
        body.height = int(prs.slide_height * 0.72)
        return slide
    rows, cols = len(spec["table"]), len(spec["table"][0])
    left = int(prs.slide_width * 0.30)
    top = int(prs.slide_height * 0.66)
    tbl = slide.shapes.add_table(rows, cols, left, top, int(prs.slide_width * 0.40), int(prs.slide_height * 0.24)).table
    for r, row in enumerate(spec["table"]):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            for para in cell.text_frame.paragraphs:
                para.alignment = 2                                  # 가운데
                for run in para.runs:
                    run.font.size = Pt(16)
                    run.font.bold = (r == 0 or c == 0)
    return slide


def video_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == 16]


def media_rid(pic_element):
    vf = pic_element.find(f".//{A}videoFile")
    return vf.get(f"{R}link") or vf.get(f"{R}embed")


def build(job, out_dir):
    src, dst = os.path.join(DECKS, job["src"]), os.path.join(DECKS, job["dst"])
    prs = Presentation(src)
    n = len(prs.slides)

    for no, name in job.get("replace", {}).items():
        mp4, png = poster(name, out_dir)
        slide = prs.slides[no - 1]
        pic = video_shapes(slide)[0]._element
        blip = pic.find(f".//{A}blip")
        slide.part.related_part(media_rid(pic))._blob = open(mp4, "rb").read()
        slide.part.related_part(blip.get(f"{R}embed"))._blob = open(png, "rb").read()
        print(f"  교체  {no:2d}번 슬라이드 ← {name}")

    added = []
    extras = []                                                     # (원본 N번 뒤, 슬라이드 색인, key)
    for after, key in job.get("extra", []):
        extra_slide(prs, key)
        extras.append((after, len(prs.slides) - 1, key))
        print(f"  문제  {after:2d}번 뒤 ← {key}")
    for after, names in job.get("insert", []):
        for name in names:
            mp4, png = poster(name, out_dir)
            slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
            slide.shapes.add_movie(mp4, 0, 0, prs.slide_width, prs.slide_height,
                                   poster_frame_image=png, mime_type="video/mp4")
            added.append((after, len(prs.slides) - 1, name))
            print(f"  추가  {after:2d}번 뒤 ← {name}")

    moved = {a - 1: b - 1 for a, b in job.get("move", [])}
    dropped = {d - 1 for d in job.get("drop", [])}       # 원본에서 빼는 장. 그 뒤에 붙는 영상은 그대로 붙는다
    order = []
    for i in range(n):
        if i in moved:
            continue
        if i not in dropped:
            order.append(i)
        else:
            print(f"  삭제  {i + 1:2d}번 슬라이드")
        order += [a for a, b in moved.items() if b == i]
        order += [idx for aft, idx, _ in extras if aft == i + 1]     # 문제 슬라이드가 그 영상보다 앞
        order += [idx for aft, idx, _ in added if aft == i + 1]

    lst = prs.slides._sldIdLst
    ids = list(lst)
    for e in ids:
        lst.remove(e)
    for i in order:
        lst.append(ids[i])
    for i in dropped:
        prs.part.drop_rel(ids[i].rId)                  # 목록에서만 빼면 파일에 남는다

    # 동기화 폴더 밖에서 저장하고 글꼴을 고친 뒤 한 번에 들여놓는다. 만든 파일을 제자리에서 다시
    # 고치면 Synology 가 "<이름> 2.pptx" 사본을 만들고 원본이 사라질 수 있다(2026-09-08).
    staged = os.path.join(out_dir, os.path.basename(dst))
    prs.save(staged)
    fix_fonts(staged, backup=False)
    shutil.copyfile(staged, dst)

    # 만들어진 순서를 그대로 적어 둔다: 원본 슬라이드는 번호, 새 영상은 씬 이름.
    by_idx = {idx: name for _, idx, name in added}
    by_extra = {idx: key for _, idx, key in extras}
    layout = [("video", by_idx[i]) if i in by_idx else (("extra", by_extra[i]) if i in by_extra else ("slide", i + 1))
              for i in order]
    for no, name in job.get("replace", {}).items():
        layout[order.index(no - 1)] = ("video", name)
    return dst, layout


def first_line(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().splitlines()[0][:46]
    return ""


def verify(job, dst, layout):
    """영상 바이트가 원본 mp4 와 같은지, 전체 화면인지, 그리고 개념 슬라이드에 맞붙어
    있는지 확인한다. 셋 중 하나라도 어긋나면 실패로 본다.

    place="before" 면 영상 **다음** 슬라이드가, "after" 면 영상 **앞** 슬라이드가
    anchors 에 적힌 번호여야 한다."""
    by_hash = {}
    for f in os.listdir(VIDEOS):
        if f.endswith(".mp4"):
            with open(os.path.join(VIDEOS, f), "rb") as fh:
                by_hash[hashlib.sha256(fh.read()).hexdigest()] = f[:-4]

    src_prs = Presentation(os.path.join(DECKS, job["src"]))
    anchors = job.get("anchors", {})
    prs = Presentation(dst)
    rows, ok = [], True

    place = job.get("place", "before")
    # 각 자리에서 바로 앞·바로 뒤의 원본 슬라이드 번호를 미리 구해 둔다.
    prev_of, next_of = [None] * len(layout), [None] * len(layout)
    seen = None
    for i, (kind, val) in enumerate(layout):
        prev_of[i] = seen
        if kind == "slide":
            seen = val
    seen = None
    for i in range(len(layout) - 1, -1, -1):
        next_of[i] = seen
        if layout[i][0] == "slide":
            seen = layout[i][1]

    for i, ((kind, val), slide) in enumerate(zip(layout, prs.slides)):
        if kind in ("slide", "extra"):
            continue
        last_slide = prev_of[i]
        pics = video_shapes(slide)
        if not pics:
            rows.append((val, "★ 영상이 없는 슬라이드", "", ""))
            ok = False
            continue
        blob = slide.part.related_part(media_rid(pics[0]._element)).blob
        name = by_hash.get(hashlib.sha256(blob).hexdigest())
        full = (abs(pics[0].width - prs.slide_width) < 1000
                and abs(pics[0].height - prs.slide_height) < 1000)
        want = anchors.get(val)
        붙은쪽 = next_of[i] if place == "before" else last_slide
        placed = (want is None or 붙은쪽 == want)
        note = ""
        if 붙은쪽 is not None:
            note = "%2d. %s" % (붙은쪽, first_line(src_prs.slides[붙은쪽 - 1]))
        if not placed:
            note += "  ★ %d번 %s 여야 한다" % (want, "앞" if place == "before" else "뒤")
        ok &= (name == val) and full and placed
        rows.append((val, "일치" if name == val else "★ mp4 불일치",
                     "전체화면" if full else "★크기이상", note))
    return rows, ok


def main():
    if not os.path.isdir(VIDEOS):
        sys.exit(f"영상 폴더가 없다: {VIDEOS}")
    only = [a for a in sys.argv[1:] if not a.startswith("-")]      # 원본 파일 이름을 주면 그 job 만 돌린다
    jobs = [j for j in JOBS if not only or j["src"] in only]
    if not jobs:
        sys.exit(f"맞는 job 이 없다: {only}. 원본 이름은 {[j['src'] for j in JOBS]}")
    out_dir = tempfile.mkdtemp(prefix="poster_")
    all_ok = True
    for job in jobs:
        print(f"═══ {job['src']} ═══")
        dst, layout = build(job, out_dir)
        rows, ok = verify(job, dst, layout)
        all_ok &= ok
        print(f"  → {os.path.basename(dst)}  {len(layout)}장 / 영상 {len(rows)}편")
        곁 = "바로 뒤 개념 슬라이드" if job.get("place", "before") == "before" else "바로 앞 개념 슬라이드"
        print(f"     {'영상':24s} {'원본대조':10s} {'크기':8s} {곁}")
        for name, match, size, note in rows:
            print(f"     {name:24s} {match:10s} {size:8s} {note}")
        print(f"  {'대조·배치 모두 통과' if ok else '★ 확인 필요'}\n")
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
